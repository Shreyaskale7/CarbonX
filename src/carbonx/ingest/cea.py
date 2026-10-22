"""Parse CEA CO2 Baseline Database workbooks (v1.0-v22.0) into factor rows.

One parser, not one per version: every release has a Results sheet whose factor
tables are a header row of FY labels ('2023-24') under or beside a title. Two
shapes occur, told apart by what the row labels below the header are:

  metric rows (v11+):  title 'Emission Factors (tCO2/MWh) (incl. Imports)',
                       rows 'Combined Margin', 'Build Margin', ...  -> India
  region rows (v1-10): title 'Combined Margin (tCO2/MWh) (incl. Imports)',
                       rows 'North', 'South', 'NEWNE', 'India', ...

Metric identity across versions is a judgment, because CEA renamed factors.
Every label is mapped by METRICS below, and an unmapped label is an error, so
a new name can never be silently dropped or merged.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path

import openpyxl
import xlrd
from openpyxl.utils import get_column_letter

# First match wins. Order matters: the specific RE/captive names before plain
# "weighted average". Evidence for the captive mapping: v19 'Grid Emission
# Factor (for net effective injection into grid)' FY2022-23 = 0.716 equals v20
# 'Weighted Average Emission Factor (for net effective injection ...)' and v22
# 'Weighted Average Grid Emission Rate (Incl. RES,Captive)' (incl. imports).
METRICS: list[tuple[str, str]] = [
    (r"net effective injection|incl\.? ?res,? ?captive", "weighted_avg_incl_re_captive"),
    (r"incl\.? ?res\b", "weighted_avg_incl_re"),
    (r"weighted average", "weighted_avg"),
    (r"operating margin", "operating_margin"),
    (r"build margin", "build_margin"),
    (r"combined margin", "combined_margin"),
]

REGIONS = {
    "india": "IN", "north": "IN-N", "east": "IN-E", "south": "IN-S",
    "west": "IN-W", "north-east": "IN-NE", "newne": "IN-NEWNE",
}

FY = re.compile(r"^\s*((?:19|20)\d\d)-(\d\d)\s*$")


@dataclass(frozen=True)
class Factor:
    version_label: str
    fy: str             # 'FY2023-24'
    factor_key: str     # 'grid.combined_margin.incl_imports'
    geography: str      # 'IN', 'IN-S', ...
    value: float
    cell: str           # 'Results!F13' -- provenance down to the cell
    label: str          # CEA's own row label, verbatim


@dataclass(frozen=True)
class Workbook:
    version_label: str
    stated_date: date | None
    date_precision: str  # 'day' | 'month' | 'unknown'
    covers_fy: str       # newest FY in the factor tables
    factors: list[Factor]


def _grid(path: Path) -> tuple[str, list[list[object]]]:
    if path.suffix == ".xls":
        book = xlrd.open_workbook(str(path))
        name = next(n for n in book.sheet_names() if n.startswith("Result"))
        sheet = book.sheet_by_name(name)
        return name, [sheet.row_values(i) for i in range(sheet.nrows)]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    name = next(n for n in wb.sheetnames if n.startswith("Result"))
    ws = wb[name]
    ws.reset_dimensions()  # some files declare a wrong used range
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return name, rows


def _cell(rows: list[list[object]], i: int, j: int) -> object:
    return rows[i][j] if i < len(rows) and j < len(rows[i]) else None


def _text(v: object) -> str:
    return " ".join(str(v).split()) if isinstance(v, str) else ""


def metric(label: str) -> str:
    low = label.lower()
    for pattern, key in METRICS:
        if re.search(pattern, low):
            return key
    raise ValueError(f"unmapped CEA factor label: {label!r}")


def imports(title: str, label: str) -> str:
    """CEA publishes each factor excluding and including inter-grid imports."""
    text = f"{title} {label}".lower()
    if "not adjusted for imports" in text:
        return "excl_imports"  # CEA's own note: BM is never import-adjusted
    if "incl. imports" in text or "(2)" in label:
        return "incl_imports"
    if "excl. imports" in text:
        return "excl_imports"
    raise ValueError(f"cannot tell import basis: {title!r} / {label!r}")


def _fy(v: object) -> str | None:
    m = FY.match(v) if isinstance(v, str) else None
    return f"FY{m[1]}-{m[2]}" if m else None


def _stated_date(rows: list[list[object]]) -> tuple[date | None, str]:
    for i in range(min(10, len(rows))):
        for j, v in enumerate(rows[i]):
            if _text(v).upper() != "DATE":
                continue
            raw = next((x for x in rows[i][j + 1:] if x not in (None, "")), None)
            if isinstance(raw, datetime):
                return raw.date(), "day"
            if isinstance(raw, float):  # Excel serial date from .xls
                return xlrd.xldate_as_datetime(raw, 0).date(), "day"
            s = _text(raw)
            for fmt, precision in (("%d %b %Y", "day"), ("%d %B %Y", "day"),
                                   ("%b %y", "month"), ("%B %y", "month")):
                try:
                    return datetime.strptime(s, fmt).replace(tzinfo=UTC).date(), precision
                except ValueError:
                    pass
            return None, "unknown"
    return None, "unknown"


def parse(path: Path, version_label: str) -> Workbook:
    sheet, rows = _grid(path)
    found: dict[tuple[str, str, str], Factor] = {}

    for i, row in enumerate(rows):
        fy_cols = [(j, fy) for j, v in enumerate(row) if (fy := _fy(v))]
        if not fy_cols:
            continue
        # Split the header into blocks at label columns: a block's label column
        # is the nearest text cell left of its first FY, in this row (v11+) or
        # the row above (v1-10).
        blocks: list[tuple[int, str, list[tuple[int, str]]]] = []
        for j, fy in fy_cols:
            if blocks and j - blocks[-1][2][-1][0] <= 1:
                blocks[-1][2].append((j, fy))
                continue
            for r in (i, i - 1):
                label_col = next((c for c in range(j - 1, -1, -1)
                                  if _text(_cell(rows, r, c))), None)
                if label_col is not None and (
                    not blocks or label_col > blocks[-1][2][-1][0]
                ):
                    break
            else:
                label_col = None
            title = "" if label_col is None else (
                _text(_cell(rows, i, label_col))
                or _text(_cell(rows, i - 1, label_col)))
            blocks.append((label_col or 0, title, [(j, fy)]))

        for label_col, title, cols in blocks:
            # Generation and emission-total tables share the FY header shape;
            # only factor tables carry the unit in their title.
            if "tco2/mwh" not in title.lower():
                continue
            for k in range(i + 1, len(rows)):
                label = _text(_cell(rows, k, label_col))
                if not label or label.startswith("("):
                    break
                region = REGIONS.get(label.lower())
                name = title if region else label
                key = f"grid.{metric(name)}.{imports(title, label)}"
                geo = region or "IN"
                for j, fy in cols:
                    v = _cell(rows, k, j)
                    if isinstance(v, bool) or not isinstance(v, (int, float)):
                        continue
                    f = Factor(version_label, fy, key, geo, float(v),
                               f"{sheet}!{get_column_letter(j + 1)}{k + 1}", label)
                    prev = found.setdefault((key, geo, fy), f)
                    if prev.value != f.value:
                        raise ValueError(f"{path.name}: {key} {geo} {fy} is both "
                                         f"{prev.value} ({prev.cell}) and "
                                         f"{f.value} ({f.cell})")

    if not found:
        raise ValueError(f"{path.name}: no factor tables found")
    stated, precision = _stated_date(rows)
    factors = sorted(found.values(), key=lambda f: (f.factor_key, f.geography, f.fy))
    return Workbook(version_label, stated, precision,
                    max(f.fy for f in factors), factors)
